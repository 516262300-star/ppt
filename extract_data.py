from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def normalize_model(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("－", "-").replace("—", "-").replace("–", "-")
    return re.sub(r"\s+", "", text).lower()


def split_source_model(value: object) -> tuple[str, list[str]]:
    text = str(value).strip() if value is not None else ""
    if not text or normalize_model(text) == "型号":
        return "", []
    parts = [part.strip() for part in re.split(r"[-－—–]", text) if part.strip()]
    return (parts[0], parts[1:]) if parts else ("", [])


def designer_name(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s*\d+(?:\.\d+)?%\s*$", "", str(value).strip())


def as_decimal(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").replace("￥", "").replace("¥", ""))
    except InvalidOperation:
        return Decimal("0")


def normalize_header(value: object) -> str:
    return re.sub(r"[\s_（）()]+", "", str(value or "").strip()).lower()


def find_header_row(sheet, required_aliases: dict[str, set[str]]) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True),
        start=1,
    ):
        headers = [normalize_header(value) for value in row]
        columns: dict[str, int] = {}
        for semantic, aliases in required_aliases.items():
            for index, header in enumerate(headers):
                if header in aliases:
                    columns[semantic] = index
                    break
        if len(columns) == len(required_aliases):
            return row_number, columns
    expected = "、".join(required_aliases)
    raise ValueError(f"工作表 {sheet.title} 前20行未找到必需列：{expected}")


def main() -> None:
    parser = argparse.ArgumentParser(description="提取设计师型号交易额与净利数据")
    parser.add_argument("--designer-file", required=True, help="分红型号/型号归属表")
    parser.add_argument("--sales-file", required=True, help="销售明细表")
    parser.add_argument("--output", required=True, help="输出 JSON")
    parser.add_argument("--designer-sheet", default="Sheet1")
    parser.add_argument("--sales-sheet", default="Worksheet")
    parser.add_argument("--period", default="")
    args = parser.parse_args()

    designer_path = Path(args.designer_file).resolve()
    sales_path = Path(args.sales_file).resolve()
    if not designer_path.exists():
        raise FileNotFoundError(f"型号归属表不存在：{designer_path}")
    if not sales_path.exists():
        raise FileNotFoundError(f"销售明细不存在：{sales_path}")

    designers_to_groups: dict[str, dict[str, dict[str, object]]] = defaultdict(dict)
    designer_order: list[str] = []
    source_rows: list[dict[str, str]] = []

    designer_wb = load_workbook(designer_path, read_only=True, data_only=True)
    try:
        if args.designer_sheet not in designer_wb.sheetnames:
            raise ValueError(
                f"型号归属工作表不存在：{args.designer_sheet}；可用：{', '.join(designer_wb.sheetnames)}"
            )
        sheet = designer_wb[args.designer_sheet]
        header_row, columns = find_header_row(
            sheet,
            {
                "model": {"型号", "model", "sku", "商品型号", "产品型号"},
                "designer": {"设计师", "负责人", "美工", "designer", "owner"},
            },
        )
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            model_value = row[columns["model"]] if columns["model"] < len(row) else None
            designer_value = row[columns["designer"]] if columns["designer"] < len(row) else None
            model_text = str(model_value).strip() if model_value is not None else ""
            name = designer_name(designer_value)
            if not model_text or not name:
                continue
            code, variants = split_source_model(model_text)
            if not code:
                continue
            if name not in designer_order:
                designer_order.append(name)
            group = designers_to_groups[name].setdefault(
                code,
                {"code": code, "variants": [], "source_models": []},
            )
            for variant in variants:
                if variant not in group["variants"]:
                    group["variants"].append(variant)
            if model_text not in group["source_models"]:
                group["source_models"].append(model_text)
            source_rows.append({"model": model_text, "designer": name})
    finally:
        designer_wb.close()

    model_lookup: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for name, groups in designers_to_groups.items():
        for code, group in groups.items():
            variants = group["variants"]
            if variants:
                for variant in variants:
                    model_lookup[normalize_model(f"{code}-{variant}")].add((name, code))
            else:
                model_lookup[normalize_model(code)].add((name, code))

    sales_totals = defaultdict(
        lambda: {"transaction": Decimal("0"), "net_profit": Decimal("0"), "rows": 0}
    )
    unmatched_sales_models: set[str] = set()

    sales_wb = load_workbook(sales_path, read_only=True, data_only=True)
    try:
        if args.sales_sheet not in sales_wb.sheetnames:
            raise ValueError(
                f"销售工作表不存在：{args.sales_sheet}；可用：{', '.join(sales_wb.sheetnames)}"
            )
        sheet = sales_wb[args.sales_sheet]
        header_row, columns = find_header_row(
            sheet,
            {
                "model": {"型号", "model", "sku", "商品型号", "产品型号"},
                "transaction": {"有效销售", "销售额", "交易金额", "成交金额", "sales", "revenue"},
                "net_profit": {"净利", "净利润", "利润", "netprofit", "profit"},
            },
        )
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            key = normalize_model(row[columns["model"]] if columns["model"] < len(row) else None)
            if not key:
                continue
            matches = model_lookup.get(key)
            if not matches:
                unmatched_sales_models.add(key)
                continue
            transaction = as_decimal(row[columns["transaction"]] if columns["transaction"] < len(row) else 0)
            net_profit = as_decimal(row[columns["net_profit"]] if columns["net_profit"] < len(row) else 0)
            for name, code in matches:
                aggregate = sales_totals[(name, code)]
                aggregate["transaction"] += transaction
                aggregate["net_profit"] += net_profit
                aggregate["rows"] += 1
    finally:
        sales_wb.close()

    payload: dict[str, object] = {
        "metadata": {
            "period": args.period,
            "designer_file": str(designer_path),
            "designer_sheet": args.designer_sheet,
            "sales_file": str(sales_path),
            "sales_sheet": args.sales_sheet,
        },
        "designers": {},
        "designer_order": designer_order,
        "source_rows": source_rows,
        "unmatched_sales_models": sorted(unmatched_sales_models)[:500],
    }

    for name in designer_order:
        rows: list[dict[str, object]] = []
        for code, group in designers_to_groups[name].items():
            aggregate = sales_totals[(name, code)]
            variants = list(group["variants"])
            display_model = code if not variants else f"{code}｜{' / '.join(variants)}"
            rows.append(
                {
                    "model": display_model,
                    "model_code": code,
                    "hole_spacings": variants,
                    "source_models": list(group["source_models"]),
                    "transaction": float(aggregate["transaction"]),
                    "net_profit": float(aggregate["net_profit"]),
                    "rows": aggregate["rows"],
                }
            )
        rows.sort(key=lambda item: (-float(item["transaction"]), str(item["model"])))
        payload["designers"][name] = {
            "model_count": len(rows),
            "matched_model_count": sum(1 for row in rows if int(row["rows"]) > 0),
            "transaction_total": round(sum(float(row["transaction"]) for row in rows), 2),
            "net_profit_total": round(sum(float(row["net_profit"]) for row in rows), 2),
            "models": rows,
        }

    output_path = Path(args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("设计师：", "、".join(designer_order))
    for name in designer_order:
        summary = payload["designers"][name]
        print(
            f"{name}: 型号 {summary['model_count']}，匹配 {summary['matched_model_count']}，"
            f"交易额 {summary['transaction_total']:,.2f}，净利 {summary['net_profit_total']:,.2f}"
        )


if __name__ == "__main__":
    main()
